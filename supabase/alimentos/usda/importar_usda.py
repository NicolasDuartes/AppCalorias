"""
importar_usda.py — primera carga de alimentos desde USDA FoodData Central.

Toma la lista curada (mapeo_usda.csv: nombre en español, categoría, cocción
y la descripción exacta del alimento en USDA), busca cada alimento en las
bases descargadas de USDA y llena una planilla con el mismo formato que
planilla_alimentos.xlsx, con los valores cada 100 g.

La planilla resultante (planilla_usda.xlsx) trae además la hoja
"Mapeo USDA": qué alimento de USDA se usó para cada uno, y cuáles no se
encontraron o necesitan revisión. Después sigue el camino de siempre:
    python ../planilla_a_sql.py --prueba planilla_usda.xlsx

Antes de correrlo, descargar de https://fdc.nal.usda.gov/download-datasets
los archivos JSON de "SR Legacy" y "Foundation Foods" y dejarlos (en .zip
o descomprimidos) en la carpeta usda/datos/.

Uso:  python importar_usda.py [carpeta_datos] [salida.xlsx]

Para encontrar la descripción exacta de un alimento en USDA (y corregir
mapeo_usda.csv), buscar por palabras en inglés:
      python importar_usda.py --buscar chicken breast raw
Muestra cada alimento que tiene todas esas palabras, con su id, sus
valores cada 100 g y sus porciones.
Necesita: pip install openpyxl
"""
import csv
import json
import re
import sys
import zipfile
from pathlib import Path

from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from columnas import CATEGORIAS, COLUMNAS  # noqa: E402
from crear_planilla import crear_libro  # noqa: E402

AQUI = Path(__file__).resolve().parent
MAPEO = AQUI / 'mapeo_usda.csv'

# Ids de nutrientes de FoodData Central
KCAL = 1008
KCAL_ATWATER = (2048, 2047)       # energía calculada (algunos alimentos de Foundation solo traen esta)
KJ = 1062
PROTEINAS = 1003
GRASAS = 1004
CARBOHIDRATOS = (1005, 1050)      # por diferencia; si falta, por suma

# Si la unidad no dice qué porción buscar, se prueba en este orden
PORCIONES_POR_DEFECTO = ['medium', 'fruit', 'whole', 'large', 'small']

AMARILLO = PatternFill('solid', start_color='FBEFC9')


def normalizar(texto):
    """Minúsculas y espacios simples. Ignora la aclaración "(Includes foods for USDA's Food
    Distribution Program)", que USDA agrega a algunas descripciones."""
    texto = re.sub(r"\(includes foods for usda's food distribution program\)", '', (texto or '').lower())
    return re.sub(r'\s+', ' ', texto).strip()


# ---------- Lectura de las bases de USDA ----------
def es_alimento(dato):
    return isinstance(dato, dict) and 'description' in dato and ('foodNutrients' in dato or 'fdcId' in dato)


def buscar_alimentos(dato):
    """Junta los alimentos de un JSON de FoodData Central, estén donde estén
    (SR Legacy los trae en "SRLegacyFoods", Foundation en "FoundationFoods", etc.)."""
    if es_alimento(dato):
        return [dato]
    encontrados = []
    if isinstance(dato, list):
        for item in dato:
            encontrados += buscar_alimentos(item)
    elif isinstance(dato, dict):
        for valor in dato.values():
            if isinstance(valor, (list, dict)):
                encontrados += buscar_alimentos(valor)
    return encontrados


def cargar_bases(carpeta):
    """Lee todos los .zip y .json de la carpeta y muestra cuántos alimentos trae cada uno."""
    alimentos = []

    def leer(nombre, contenido):
        try:
            encontrados = buscar_alimentos(json.loads(contenido))
        except ValueError:
            print(f'  {nombre}: no es un JSON válido, se saltea')
            return
        print(f'  {nombre}: {len(encontrados)} alimentos')
        alimentos.extend(encontrados)

    for ruta in sorted(carpeta.glob('*')):
        if ruta.suffix.lower() == '.zip':
            with zipfile.ZipFile(ruta) as archivo:
                jsons = [n for n in archivo.namelist() if n.lower().endswith('.json')]
                if not jsons:
                    print(f'  {ruta.name}: no trae archivos .json (¿se descargó la versión CSV?)')
                for nombre in jsons:
                    leer(f'{ruta.name} → {nombre}', archivo.read(nombre))
        elif ruta.suffix.lower() == '.json':
            leer(ruta.name, ruta.read_bytes())
    return alimentos


def nutrientes(alimento):
    """{id de nutriente: cantidad cada 100 g}"""
    valores = {}
    for item in alimento.get('foodNutrients', []):
        nutriente = item.get('nutrient', {})
        id_nutriente = nutriente.get('id', item.get('nutrientId'))
        cantidad = item.get('amount', item.get('value'))
        if id_nutriente is not None and cantidad is not None:
            valores[int(id_nutriente)] = float(cantidad)
    return valores


def valores_cada_100g(alimento):
    """kcal, proteínas, carbohidratos y grasas, o None si falta alguno."""
    n = nutrientes(alimento)
    kcal = n.get(KCAL)
    if kcal is None:
        kcal = next((n[i] for i in KCAL_ATWATER if i in n), None)
    if kcal is None and KJ in n:
        kcal = n[KJ] / 4.184
    carbohidratos = next((n[i] for i in CARBOHIDRATOS if i in n), None)
    valores = {'kcal': kcal, 'proteinas': n.get(PROTEINAS), 'carbohidratos': carbohidratos, 'grasas': n.get(GRASAS)}
    if None in valores.values():
        return None
    # Algunos alimentos traen carbohidratos por diferencia apenas negativos (redondeo de USDA)
    return {clave: round(max(valor, 0), 1) for clave, valor in valores.items()}


def gramos_por_unidad(alimento, buscar):
    """Peso de 1 unidad según las porciones de USDA ("1 large = 50 g"), o None."""
    porciones = alimento.get('foodPortions', [])
    palabras = [buscar] if buscar else PORCIONES_POR_DEFECTO
    for palabra in palabras:
        for porcion in porciones:
            texto = ' '.join(str(porcion.get(c) or '') for c in ('modifier', 'portionDescription'))
            texto += ' ' + str((porcion.get('measureUnit') or {}).get('name') or '')
            if re.search(rf'\b{re.escape(palabra)}\b', texto.lower()) and porcion.get('gramWeight'):
                cantidad = float(porcion.get('amount') or 1) or 1
                return round(float(porcion['gramWeight']) / cantidad, 1)
    return None


def candidatos(descripcion, por_descripcion, cantidad=3):
    """Los alimentos de USDA con la descripción más parecida (para corregir el mapeo)."""
    buscadas = set(re.findall(r'[a-z0-9%]+', normalizar(descripcion)))
    primera = normalizar(descripcion).split(',')[0]
    puntajes = []
    for clave, alimento in por_descripcion.items():
        if not clave.startswith(primera):
            continue
        palabras = set(re.findall(r'[a-z0-9%]+', clave))
        puntaje = len(buscadas & palabras) / len(buscadas | palabras)
        puntajes.append((puntaje, alimento['description']))
    return [d for _, d in sorted(puntajes, reverse=True)[:cantidad]]


# ---------- Mapeo ----------
def leer_mapeo(ruta):
    # Excel en español guarda los CSV en UTF-8 o en Windows-1252 según cómo se elija guardar
    for codificacion in ('utf-8-sig', 'cp1252'):
        try:
            with open(ruta, encoding=codificacion, newline='') as archivo:
                return list(csv.DictReader(archivo, delimiter=';'))
        except UnicodeDecodeError:
            continue
    raise SystemExit(f'No se pudo leer {ruta.name}: guardalo desde Excel como "CSV UTF-8 (delimitado por comas)".')


def mapear(fila, por_descripcion):
    """Devuelve (datos para la planilla o None, estado, detalle)."""
    alimento = por_descripcion.get(normalizar(fila['usda_descripcion']))
    if alimento is None:
        return None, 'No encontrado', 'Parecidos en USDA: ' + ' | '.join(candidatos(fila['usda_descripcion'], por_descripcion))

    valores = valores_cada_100g(alimento)
    if valores is None:
        return None, 'Faltan nutrientes', 'USDA no trae kcal, proteínas, carbohidratos o grasas para este alimento'

    datos = {
        'nombre': fila['nombre'].strip(),
        'coccion': fila['coccion'].strip(),
        **valores,
        'sinonimos': fila['sinonimos'].strip(),
        'fuente': 'usda',
        'fuente_ref': str(alimento.get('fdcId', '')),
    }
    estado, detalle = 'OK', ''
    if fila['unidad_nombre'].strip():
        gramos = gramos_por_unidad(alimento, fila['unidad_buscar'].strip())
        if gramos:
            datos['unidad_nombre'] = fila['unidad_nombre'].strip()
            datos['unidad_g'] = gramos
            datos['porcion_g'] = gramos
            datos['porcion_texto'] = f'1 {fila["unidad_nombre"].strip()}'
        else:
            estado, detalle = 'Revisar', 'No se encontró el peso de 1 unidad en las porciones de USDA'
    return datos, estado, detalle


# ---------- Salida ----------
def escribir_fila(hoja, datos):
    claves = [c['clave'] for c in COLUMNAS]
    fila = next(f for f in range(2, hoja.max_row + 2) if not hoja.cell(row=f, column=1).value)
    for clave, valor in datos.items():
        hoja.cell(row=fila, column=claves.index(clave) + 1, value=valor)


def hoja_mapeo(libro, informe):
    hoja = libro.create_sheet('Mapeo USDA', 1)
    encabezado = ['nombre', 'coccion', 'categoria', 'estado', 'usda_descripcion', 'fdcId', 'detalle', 'nota']
    hoja.append(encabezado)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in informe:
        hoja.append([fila.get(c, '') for c in encabezado])
        if fila['estado'] != 'OK':
            for celda in hoja[hoja.max_row]:
                celda.fill = AMARILLO
    for letra, ancho in zip('ABCDEFGH', (30, 9, 14, 17, 60, 10, 70, 50)):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = 'A2'


def buscar(palabras, carpeta, maximo=40):
    """Lista los alimentos de USDA cuya descripción tiene todas las palabras."""
    palabras = [p.lower() for p in palabras]
    alimentos = cargar_bases(carpeta)
    encontrados = [a for a in alimentos if all(p in a['description'].lower() for p in palabras)]
    encontrados.sort(key=lambda a: (len(a['description']), a['description']))
    print(f'\n{len(encontrados)} alimentos con "{" ".join(palabras)}"' + (f' (se muestran {maximo})' if len(encontrados) > maximo else '') + ':\n')
    for alimento in encontrados[:maximo]:
        valores = valores_cada_100g(alimento)
        texto = (f"{valores['kcal']} kcal · P {valores['proteinas']} · C {valores['carbohidratos']} · G {valores['grasas']}"
                 if valores else 'faltan valores')
        porciones = '; '.join(
            f"{p.get('amount') or 1} {p.get('modifier') or p.get('portionDescription') or (p.get('measureUnit') or {}).get('name', '')}"
            f" = {p.get('gramWeight')} g"
            for p in alimento.get('foodPortions', [])[:4])
        print(f"{alimento['description']}\n    id {alimento.get('fdcId')} · {texto}" + (f'\n    porciones: {porciones}' if porciones else ''))


def main():
    if len(sys.argv) > 1 and sys.argv[1] == '--buscar':
        buscar(sys.argv[2:], AQUI / 'datos')
        return
    carpeta = Path(sys.argv[1]) if len(sys.argv) > 1 else AQUI / 'datos'
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else AQUI / 'planilla_usda.xlsx'

    print(f'Leyendo las bases de USDA en {carpeta}:')
    alimentos_usda = cargar_bases(carpeta)
    if not alimentos_usda:
        print(f'No hay bases de USDA en {carpeta}. Descargá los JSON de "SR Legacy" y "Foundation Foods"')
        print('desde https://fdc.nal.usda.gov/download-datasets y dejalos en esa carpeta.')
        sys.exit(1)
    # Si una descripción está en las dos bases, se usa la que tenga los 4 valores completos
    por_descripcion = {}
    for alimento in alimentos_usda:
        clave = normalizar(alimento.get('description'))
        if clave not in por_descripcion or (valores_cada_100g(por_descripcion[clave]) is None
                                            and valores_cada_100g(alimento) is not None):
            por_descripcion[clave] = alimento

    libro = crear_libro()
    informe, cargados = [], 0
    for fila in leer_mapeo(MAPEO):
        categoria = fila['categoria'].strip()
        if categoria not in CATEGORIAS:
            informe.append({**fila, 'estado': 'Revisar', 'detalle': f'La categoría "{categoria}" no existe'})
            continue
        datos, estado, detalle = mapear(fila, por_descripcion)
        alimento = por_descripcion.get(normalizar(fila['usda_descripcion'])) or {}
        informe.append({**fila, 'estado': estado, 'detalle': detalle, 'fdcId': alimento.get('fdcId', '')})
        if datos:
            escribir_fila(libro[CATEGORIAS[categoria]], datos)
            cargados += 1

    hoja_mapeo(libro, informe)
    libro.save(salida)

    faltan = [f for f in informe if f['estado'] != 'OK']
    print(f'{len(alimentos_usda)} alimentos leídos de USDA.')
    print(f'Listo: {cargados} de {len(informe)} alimentos en {salida.name}.')
    if faltan:
        print(f'{len(faltan)} para revisar (ver la hoja "Mapeo USDA"):')
        for f in faltan:
            print(f'  - {f["nombre"]} ({f["coccion"]}): {f["estado"]}')


if __name__ == '__main__':
    main()
