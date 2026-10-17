"""
planilla_a_sql.py — convierte la planilla de alimentos en la carga para Supabase.

Lee planilla_alimentos.xlsx (una hoja por categoría), revisa cada fila con
los mismos controles que la base y genera carga_alimentos.sql. Ese archivo
se pega en Supabase → SQL Editor → Run. Se puede correr las veces que haga
falta: si un alimento ya existe (mismo nombre, marca y cocción, sin importar
tildes ni mayúsculas) se actualiza en vez de duplicarse.

Si alguna fila tiene errores, no genera nada y dice en qué hoja y fila están.

Uso:  python planilla_a_sql.py [planilla.xlsx] [salida.sql]
Necesita: pip install openpyxl
"""
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

from columnas import CATEGORIAS, COCCIONES, COLUMNAS, FUENTES

AQUI = Path(__file__).parent
CLAVES = [c['clave'] for c in COLUMNAS]
FUENTES_PLANILLA = [f for f in FUENTES if f not in ('coach', 'alumno')]


def normalizar(texto):
    """Igual que normalizar_texto() de la base: minúsculas, sin tildes ni espacios de más."""
    sin_tildes = ''.join(c for c in unicodedata.normalize('NFD', texto or '') if unicodedata.category(c) != 'Mn')
    return ' '.join(sin_tildes.lower().split())


def texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def numero(valor):
    """Acepta 12,5 o 12.5. Devuelve None si está vacío y ValueError si no es un número."""
    valor = texto(valor)
    if valor == '':
        return None
    return float(valor.replace(',', '.'))


def kcal_coinciden(kcal, p, c, g):
    return abs(4 * p + 4 * c + 9 * g - kcal) <= max(0.2 * kcal, 15)


def revisar_fila(datos):
    """Devuelve (alimento, errores) con los mismos controles que la migración."""
    errores = []
    alimento = {'nombre': ' '.join(texto(datos['nombre']).split())}
    if not 2 <= len(alimento['nombre']) <= 80:
        errores.append('el nombre tiene que tener entre 2 y 80 letras')

    alimento['coccion'] = normalizar(texto(datos['coccion']))
    if alimento['coccion'] not in COCCIONES:
        errores.append('cocción tiene que ser crudo, cocido o listo')

    for clave in ('kcal', 'proteinas', 'carbohidratos', 'grasas', 'porcion_g', 'unidad_g'):
        try:
            alimento[clave] = numero(datos[clave])
        except ValueError:
            errores.append(f'{clave} no es un número')
            alimento[clave] = None

    k, p, c, g = (alimento[x] for x in ('kcal', 'proteinas', 'carbohidratos', 'grasas'))
    if None in (k, p, c, g):
        if not errores:
            errores.append('faltan kcal, proteínas, carbohidratos o grasas')
    else:
        if not 0 < k <= 900:
            errores.append('kcal tiene que estar entre 0 y 900')
        if min(p, c, g) < 0 or max(p, c, g) > 100:
            errores.append('cada macro tiene que estar entre 0 y 100 g')
        if p + c + g > 100:
            errores.append('proteínas + carbohidratos + grasas pasa de 100 g')
        excepcion = normalizar(texto(datos['excepcion_kcal'])) in ('si', 'sí')
        if not excepcion and not kcal_coinciden(k, p, c, g):
            errores.append(f'las kcal no coinciden con los macros (dan {round(4 * p + 4 * c + 9 * g)}); '
                           'si es a propósito, poné "sí" en excepcion_kcal')

    for clave in ('porcion_g', 'unidad_g'):
        if alimento[clave] is not None and not 0 < alimento[clave] <= 2000:
            errores.append(f'{clave} tiene que estar entre 0 y 2000')

    alimento['porcion_texto'] = texto(datos['porcion_texto']) or None
    if alimento['porcion_texto'] and len(alimento['porcion_texto']) > 40:
        errores.append('porcion_texto tiene más de 40 letras')

    alimento['unidad_nombre'] = texto(datos['unidad_nombre']) or None
    if (alimento['unidad_nombre'] is None) != (alimento['unidad_g'] is None):
        errores.append('unidad_nombre y unidad_g van juntos (los dos o ninguno)')

    alimento['marca'] = texto(datos['marca']) or None
    if alimento['marca'] and len(alimento['marca']) > 60:
        errores.append('la marca tiene más de 60 letras')

    alimento['sinonimos'] = [s.strip() for s in texto(datos['sinonimos']).split(',') if s.strip()]

    alimento['fuente'] = normalizar(texto(datos['fuente'])) or 'equipo'
    if alimento['fuente'] not in FUENTES_PLANILLA:
        errores.append('fuente tiene que ser ' + ', '.join(FUENTES_PLANILLA))
    alimento['fuente_ref'] = texto(datos['fuente_ref']) or None
    return alimento, errores


def leer_planilla(ruta):
    libro = load_workbook(ruta, data_only=True, read_only=True)
    codigo_de_hoja = {nombre: codigo for codigo, nombre in CATEGORIAS.items()}
    alimentos, errores, claves = [], [], {}

    for hoja in libro.worksheets:
        if hoja.title == 'Instrucciones':
            continue
        if hoja.title not in codigo_de_hoja:
            errores.append(f'Hoja "{hoja.title}": no es una categoría conocida (no cambies los nombres de las hojas)')
            continue
        filas = hoja.iter_rows(values_only=True)
        encabezado = [texto(v) for v in next(filas, [])]
        faltan = [c for c in CLAVES if c != 'control' and c not in encabezado]
        if faltan:
            errores.append(f'Hoja "{hoja.title}": faltan las columnas {", ".join(faltan)}')
            continue

        for numero_fila, valores in enumerate(filas, start=2):
            datos = {clave: (valores[encabezado.index(clave)] if encabezado.index(clave) < len(valores) else None)
                     for clave in CLAVES if clave in encabezado}
            if not texto(datos.get('nombre')):
                continue  # fila vacía (las fórmulas de "control" no cuentan)
            alimento, problemas = revisar_fila(datos)
            alimento['categoria'] = codigo_de_hoja[hoja.title]
            lugar = f'Hoja "{hoja.title}", fila {numero_fila} ({alimento["nombre"]})'
            errores += [f'{lugar}: {p}' for p in problemas]

            clave = f'{normalizar(alimento["nombre"])}|{normalizar(alimento["marca"])}|{alimento["coccion"]}'
            if clave in claves:
                errores.append(f'{lugar}: repetido, ya está en {claves[clave]}')
            claves[clave] = lugar
            alimentos.append(alimento)
    return alimentos, errores


def sql_texto(valor):
    if valor is None:
        return 'null'
    return "'" + str(valor).replace("'", "''") + "'"


def sql_numero(valor):
    return 'null' if valor is None else repr(round(valor, 1))


def sql_lista(valores):
    return "'{}'::text[]" if not valores else 'array[' + ', '.join(sql_texto(v) for v in valores) + ']::text[]'


def generar_sql(alimentos):
    columnas = ['nombre', 'categoria', 'coccion', 'marca', 'kcal', 'proteinas', 'carbohidratos', 'grasas',
                'porcion_g', 'porcion_texto', 'unidad_nombre', 'unidad_g', 'sinonimos', 'fuente', 'fuente_ref']
    filas = []
    for a in alimentos:
        valores = [sql_texto(a['nombre']), sql_texto(a['categoria']), sql_texto(a['coccion']), sql_texto(a['marca']),
                   sql_numero(a['kcal']), sql_numero(a['proteinas']), sql_numero(a['carbohidratos']),
                   sql_numero(a['grasas']), sql_numero(a['porcion_g']), sql_texto(a['porcion_texto']),
                   sql_texto(a['unidad_nombre']), sql_numero(a['unidad_g']), sql_lista(a['sinonimos']),
                   sql_texto(a['fuente']), sql_texto(a['fuente_ref'])]
        filas.append('  (' + ', '.join(valores) + ", 'aprobado')")
    actualizar = ',\n  '.join(f'{c} = excluded.{c}' for c in columnas if c != 'nombre')
    return (
        '-- Generado por supabase/alimentos/planilla_a_sql.py: no editar a mano.\n'
        '-- Pegar en Supabase → SQL Editor → Run. Corre en una transacción.\n\n'
        'begin;\n\n'
        f'insert into public.alimentos ({", ".join(columnas)}, estado_revision)\nvalues\n'
        + ',\n'.join(filas) +
        "\non conflict (clave) where estado_revision <> 'rechazado' do update set\n  "
        + actualizar + ",\n  estado_revision = 'aprobado';\n\n"
        f"select count(*) as alimentos_aprobados from public.alimentos where estado_revision = 'aprobado';\n\n"
        'commit;\n'
    )


def main():
    planilla = Path(sys.argv[1]) if len(sys.argv) > 1 else AQUI / 'planilla_alimentos.xlsx'
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else AQUI / 'carga_alimentos.sql'

    alimentos, errores = leer_planilla(planilla)
    if errores:
        print(f'La planilla tiene {len(errores)} error(es). Corregilos y volvé a correr el script:\n')
        print('\n'.join(f'  - {e}' for e in errores))
        sys.exit(1)
    if not alimentos:
        print('La planilla no tiene alimentos cargados.')
        sys.exit(1)

    salida.write_text(generar_sql(alimentos), encoding='utf-8')
    por_categoria = {}
    for a in alimentos:
        por_categoria[CATEGORIAS[a['categoria']]] = por_categoria.get(CATEGORIAS[a['categoria']], 0) + 1
    print(f'Listo: {len(alimentos)} alimento(s) en {salida.name}')
    for nombre, cantidad in por_categoria.items():
        print(f'  {nombre}: {cantidad}')


if __name__ == '__main__':
    main()
