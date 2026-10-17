"""
crear_planilla.py — arma la planilla modelo para cargar alimentos.

Genera planilla_alimentos.xlsx: una hoja por categoría (las mismas de la tabla
categorias_alimento), con listas desplegables y una columna "control" que
marca en rojo los errores mientras se carga. Correrlo de nuevo pisa la
planilla: usarlo solo para empezar una vacía o si cambian las columnas.

Uso:  python crear_planilla.py
Necesita: pip install openpyxl
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from columnas import CATEGORIAS, COCCIONES, COLUMNAS, FUENTES

DESTINO = Path(__file__).with_name('planilla_alimentos.xlsx')
FILAS = 500  # filas preparadas (con fórmulas y validaciones) por hoja

ROJO = PatternFill('solid', start_color='F6D5CF')
ENCABEZADO = PatternFill('solid', start_color='E4EFE8')


def letra(clave):
    return chr(ord('A') + [c['clave'] for c in COLUMNAS].index(clave))


def formula_control(fila):
    c = {col['clave']: f'{letra(col["clave"])}{fila}' for col in COLUMNAS}
    kcal_macros = f'(4*{c["proteinas"]}+4*{c["carbohidratos"]}+9*{c["grasas"]})'
    return (
        f'=IF({c["nombre"]}="","",'
        f'IF(OR({c["coccion"]}="",{c["kcal"]}="",{c["proteinas"]}="",{c["carbohidratos"]}="",{c["grasas"]}=""),"Faltan datos",'
        f'IF({c["proteinas"]}+{c["carbohidratos"]}+{c["grasas"]}>100,"Macros suman más de 100 g",'
        f'IF(OR({c["kcal"]}<=0,{c["kcal"]}>900),"kcal fuera de rango",'
        f'IF(AND(ABS({kcal_macros}-{c["kcal"]})>MAX(0.2*{c["kcal"]},15),{c["excepcion_kcal"]}<>"sí"),'
        f'"kcal no coinciden: macros dan "&ROUND({kcal_macros},0),'
        f'IF(({c["unidad_nombre"]}="")<>({c["unidad_g"]}=""),"Unidad incompleta","OK"))))))'
    )


def hoja_instrucciones(libro):
    hoja = libro.active
    hoja.title = 'Instrucciones'
    filas = [
        ['Planilla de alimentos de AppCalorias'],
        [],
        ['Cómo cargar'],
        ['1. Cada hoja es una categoría. Cargá cada alimento en la hoja que le corresponde, una fila por alimento.'],
        ['2. Los valores van CADA 100 g (la columna "por 100 g" de la etiqueta), no por porción.'],
        ['3. La columna "control" tiene que decir OK. Si dice otra cosa (en rojo), corregí la fila.'],
        ['4. No cambies los nombres de las hojas ni de las columnas: el script los usa para leer la planilla.'],
        ['5. Para subirla: python planilla_a_sql.py → genera carga_alimentos.sql → pegarlo en Supabase (SQL Editor) y Run.'],
        [],
        ['Columnas'],
    ]
    filas += [[f'{c["clave"]}', c['ayuda']] for c in COLUMNAS]
    filas += [
        [],
        ['Ejemplos (no se cargan: son solo para ver el formato)'],
        [c['clave'] for c in COLUMNAS],
        ['Huevo', 'cocido', 155, 13, 1.1, 11, 50, '1 unidad', 'u', 50, 'huevo duro', '', 'usda', '173424', 'no'],
        ['Avena arrollada', 'crudo', 375, 12.5, 67.5, 7.5, 40, '4 cucharadas', '', '', 'copos de avena, avena', 'Quaker', 'equipo', '', 'no'],
    ]
    for fila in filas:
        hoja.append(fila)
    hoja['A1'].font = Font(bold=True, size=14)
    for celda in ('A3', 'A10'):
        hoja[celda].font = Font(bold=True)
    hoja.column_dimensions['A'].width = 24
    hoja.column_dimensions['B'].width = 90


def hoja_categoria(libro, nombre):
    hoja = libro.create_sheet(nombre)
    hoja.append([c['clave'] for c in COLUMNAS])
    for i, col in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=i)
        celda.font = Font(bold=True)
        celda.fill = ENCABEZADO
        celda.alignment = Alignment(horizontal='center')
        hoja.column_dimensions[letra(col['clave'])].width = col['ancho']
    hoja.freeze_panes = 'B2'

    ultima = FILAS + 1
    listas = [
        ('coccion', COCCIONES),
        ('fuente', [f for f in FUENTES if f not in ('coach', 'alumno')]),
        ('excepcion_kcal', ['sí', 'no']),
    ]
    for clave, opciones in listas:
        validacion = DataValidation(type='list', formula1='"' + ','.join(opciones) + '"', allow_blank=True)
        validacion.error = 'Elegí una opción de la lista.'
        hoja.add_data_validation(validacion)
        validacion.add(f'{letra(clave)}2:{letra(clave)}{ultima}')

    for clave, minimo, maximo in [('kcal', 0, 900), ('proteinas', 0, 100), ('carbohidratos', 0, 100),
                                  ('grasas', 0, 100), ('porcion_g', 0, 2000), ('unidad_g', 0, 2000)]:
        validacion = DataValidation(type='decimal', operator='between', formula1=str(minimo),
                                    formula2=str(maximo), allow_blank=True)
        validacion.error = f'Tiene que ser un número entre {minimo} y {maximo}.'
        hoja.add_data_validation(validacion)
        validacion.add(f'{letra(clave)}2:{letra(clave)}{ultima}')

    control = letra('control')
    for fila in range(2, ultima + 1):
        hoja[f'{control}{fila}'] = formula_control(fila)
    hoja.conditional_formatting.add(
        f'A2:{control}{ultima}',
        FormulaRule(formula=[f'AND(${control}2<>"",${control}2<>"OK")'], fill=ROJO))


def crear_libro():
    """La planilla vacía, lista para llenar (la usa también usda/importar_usda.py)."""
    libro = Workbook()
    hoja_instrucciones(libro)
    for nombre in CATEGORIAS.values():
        hoja_categoria(libro, nombre)
    return libro


def main():
    crear_libro().save(DESTINO)
    print(f'Planilla creada: {DESTINO}')


if __name__ == '__main__':
    main()
